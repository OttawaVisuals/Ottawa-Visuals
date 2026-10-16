#!/usr/bin/env python3
"""Pull OC Transpo KPI numbers out of the harvested eScribe report PDFs.

Reads the text/table extracts that oc_transpo_kpi_scraper.py --extract left in
OC_Transpo/data/ (gitignored, local only) and writes a tidy, committed history
to OC_Transpo/kpi_history/:

  kpi_observations.csv : every value found, one row per (source, metric,
                         period) -- the same month can appear in several
                         reports, and later editions sometimes revise it.
  kpi_series.csv       : one preferred value per metric/mode/breakdown/period
                         (latest publication wins; an older value that agrees
                         within rounding is kept when it is more precise).
  extract_log.txt      : per-document coverage, specs that matched nothing,
                         and every revision/conflict between sources.

The reports change format every couple of years, so extraction is driven by
explicit per-document specs rather than guessing:
  * scorecard   -- 5-Year Roadmap scorecards (2024+), one column per year
  * monthly     -- month-by-row tables in the 2020-2021 performance reports
  * chart_rows  -- chart data labels that survive in the text layer
                   ("2022 Actual 3.0 M 3.4 M ...") in the 2023 update deck
  * sequence    -- a run of N labels on one slide mapped to consecutive months
  * narrative   -- single figures quoted in report prose (regex must match)
Charts drawn as images (most monthly charts after 2021) can't be recovered.

Recent months come from the Open Ottawa KPI spreadsheets that snapshot_kpis.py
archives in rt_data/kpi_snapshots/ (rolling ~13-month windows, committed). Each
dated snapshot counts as its own source, so overlapping snapshots that revise a
month show up in the log like any other revision. Their measures were checked
against the scorecards: 2025 monthly averages reproduce the annual figures.

Stdlib only, plus openpyxl for the snapshots (skipped with a log line if it's
missing). Run from anywhere:  python OC_Transpo/scripts/extract_kpis.py
"""
import csv
import io
import re
import sys
from collections import defaultdict
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"
SNAPSHOTS = ROOT / "rt_data" / "kpi_snapshots"
OUT = ROOT / "kpi_history"

MONTHS = {m: i for i, m in enumerate(
    ["january", "february", "march", "april", "may", "june", "july",
     "august", "september", "october", "november", "december"], 1)}
MONTHS.update({k[:3]: v for k, v in list(MONTHS.items())})

FIELDS = ["metric", "mode", "breakdown", "period", "period_type", "value", "unit",
          "meeting_date", "source_file", "page", "method", "note"]
METHOD_RANK = {"monthly": 0, "scorecard": 0, "snapshot": 0, "chart_rows": 1, "sequence": 1, "narrative": 2}


# --------------------------------------------------------------------------- docs

class Doc:
    """One extracted PDF: page texts plus pdfplumber table blocks."""

    def __init__(self, key):
        self.key = key  # "<meeting folder>/<pdf stem>"
        self.meeting_date = key[:10]
        txt = DATA / (key + ".txt")
        tab = DATA / (key + ".tables.csv")
        self.pages = {}
        if txt.exists():
            parts = re.split(r"^----- page (\d+) -----\s*$",
                             txt.read_text(encoding="utf-8"), flags=re.M)
            for i in range(1, len(parts), 2):
                self.pages[int(parts[i])] = parts[i + 1]
        self.blocks = []  # (page, [cleaned rows])
        if tab.exists():
            parts = re.split(r"^# page (\d+)\s*$", tab.read_text(encoding="utf-8"), flags=re.M)
            for i in range(1, len(parts), 2):
                rows = [[re.sub(r"\s+", " ", c).strip() for c in r]
                        for r in csv.reader(io.StringIO(parts[i + 1].strip("\n")))]
                self.blocks.append((int(parts[i]), [r for r in rows if any(r)]))

    @property
    def exists(self):
        return bool(self.pages or self.blocks)

    def flat(self, page):
        return re.sub(r"\s+", " ", self.pages.get(page, ""))


def rec(doc, page, method, metric, mode, breakdown, period, period_type, value, unit, note="", prec=1):
    return {"_prec": prec, "metric": metric, "mode": mode, "breakdown": breakdown, "period": period,
            "period_type": period_type, "value": value, "unit": unit,
            "meeting_date": doc.meeting_date, "source_file": doc.key + ".pdf",
            "page": page, "method": method, "note": note}


# ------------------------------------------------------------------------ values

def parse_value(s):
    """'76.13%' -> 76.13, '8:16 mins' -> 8.27 (minutes), '64.2 M' -> 64200000,
    '57.0 k' -> 57000, '$11.3 M' -> 11300000, '6 weeks' -> 6, '1,343' -> 1343.
    Returns (value, note, precision) -- precision is the rounding step of the
    printed figure (0.1 M -> 100000) -- or (None, '', None) for text cells."""
    s = (s or "").strip()
    note = ""
    if "*" in s:
        note = "asterisked in source (see report footnote)"
        s = s.replace("*", "").strip()
    if not s or s.upper().startswith("N/A") or s in {"-", "–"}:
        return None, "", None
    m = re.fullmatch(r"(\d{1,2}):(\d{2})(?: ?mins?)?", s)
    if m:
        return round(int(m.group(1)) + int(m.group(2)) / 60, 2), note, 1 / 60
    m = re.fullmatch(r"\$?(-?[\d,]*\.?\d+) ?([MkK%]|weeks?)?", s)
    if not m:
        return None, "", None
    num = m.group(1).replace(",", "")
    scale = {"M": 1e6, "k": 1e3, "K": 1e3}.get(m.group(2) or "", 1)
    v = float(num) * scale
    prec = 10 ** -(len(num.split(".")[1]) if "." in num else 0) * scale
    return (int(v) if v == int(v) and scale != 1 else round(v, 4)), note, prec


def month_period(label, default_year):
    """'July 2020' / 'January' -> '2020-07'; None when the label isn't a month."""
    m = re.fullmatch(r"([A-Za-z]+)\.?(?: (\d{4}))?", label.strip())
    if not m or m.group(1).lower() not in MONTHS:
        return None
    year = int(m.group(2)) if m.group(2) else default_year
    return f"{year}-{MONTHS[m.group(1).lower()]:02d}" if year else None


def total_period(label, default_year, log, where):
    """Aggregate rows: '2020 Overall' -> ('2020','year'); 'Total' -> default year;
    'July 2020 to June 2021 Overall' -> ('2020-07/2021-06','12mo')."""
    lab = label.strip()
    m = re.fullmatch(r"(\d{4}) (?:Overall|Total)", lab)
    if m:
        return m.group(1), "year"
    if lab in {"Total", "Overall"} and default_year:
        return str(default_year), "year"
    m = re.fullmatch(r"([A-Za-z]+) (\d{4}) to ([A-Za-z]+) (\d{4})(?: Overall| Total)?", lab)
    if m:
        a = month_period(f"{m.group(1)} {m.group(2)}", None)
        b = month_period(f"{m.group(3)} {m.group(4)}", None)
        if a and b and a > b:  # the June 2021 report prints "July 2021 to June 2020"
            log.append(f"  {where}: label '{lab}' runs backwards; read as "
                       f"{m.group(1)} {m.group(4)} to {m.group(3)} {m.group(2)}")
            a, b = (month_period(f"{m.group(1)} {m.group(4)}", None),
                    month_period(f"{m.group(3)} {m.group(2)}", None))
        if a and b:
            return f"{a}/{b}", "12mo"
    return None, None


# -------------------------------------------------------------------- extractors

# Scorecard rows: regex on the row label -> (metric, mode, unit). Labels drift
# between editions ("Transit Commission" -> "Transit Committee"), hence regexes.
SCORECARD_MAP = [
    (r"hiring plan|staffing and recruitment", "hiring_plan_vs_actual", "org", "%"),
    (r"outreach initiatives", "outreach_events", "org", "count"),
    (r"time to fill bus operator", "operator_vacancy_fill_time", "org", "weeks"),
    (r"number of diversity initiatives", "diversity_initiatives", "org", "count"),
    (r"staff retention", "staff_retention", "org", "%"),
    (r"number of employee events", "employee_events", "org", "count"),
    (r"attendance at employee events", "employee_event_attendance", "org", "count"),
    (r"meetings held to review employee survey", "survey_review_meetings", "org", "count"),
    (r"workplace injury rate", "workplace_injury_rate", "org", "injuries per 200,000 hrs"),
    (r"preventable vehicle collisions", "preventable_collisions", "bus", "per 100,000 km"),
    (r"special constable response", "special_constable_response_time", "system", "min"),
    (r"customer injury rate", "customer_injury_rate", "system", "injuries per million trips"),
    (r"tsb-reportable", "tsb_reportable_occurrences", "otrain", "count"),
    (r"my para transpo usage", "para_online_booking_share", "para", "%"),
    (r"telephone booking line wait", "para_booking_call_wait", "para", "min"),
    (r"para transpo customer complaints", "para_complaints_per_1000_trips", "para", "per 1,000 trips"),
    (r"call centre service level", "call_centre_wait", "system", "min"),
    (r"digital certainty index", "digital_certainty_index", "system", "index"),
    (r"presentations to transit", "council_presentations", "org", "count"),
    (r"technical briefings", "council_technical_briefings", "org", "count"),
    (r"memos issued to council", "council_memos", "org", "count"),
    (r"individual speakers", "committee_public_speakers", "org", "count"),
    (r"consultation sessions held", "route_review_consultations", "bus", "count"),
    (r"written feedback received", "route_review_written_feedback", "bus", "count"),
    (r"^bus service delivery", "trips_delivered", "bus", "%"),
    (r"line 2 and 4 service delivery", "trips_delivered", "line2_4", "%"),
    (r"line 1 service delivery", "trips_delivered", "line1", "%"),
    (r"buses unavailable to meet service", "bus_fleet_unavailable", "bus", "%"),
    (r"bus operator availability", "bus_operator_availability", "bus", "%"),
    (r"bus type mismatch", "bus_type_match", "bus", "%"),
    (r"mean distance between failure", "bus_mean_distance_between_failure", "bus", "km"),
    (r"^total ridership", "ridership", "system", "linked trips"),
    (r"regularity of frequent routes", "regularity_frequent_routes", "bus", "%"),
    (r"punctuality of bus routes|routes operating every 16 minutes", "punctuality_less_frequent_routes", "bus", "%"),
    (r"para transpo on-time", "para_on_time", "para", "%"),
    (r"electric buses integrated", "electric_buses_in_fleet", "bus", "count"),
    (r"average distance travelled per electric", "electric_bus_avg_distance", "bus", "km"),
    (r"charger reliability", "electric_bus_charger_reliability", "bus", "%"),
    (r"documents available on sharepoint", "documents_on_sharepoint", "org", "%"),
    (r"kpis updated and reported", "kpis_reported", "org", "count"),
]


def scorecard(doc, spec, log):
    out, unmapped = [], set()
    for page, rows in doc.blocks:
        years = None
        for r in rows:
            hdr = {i: re.search(r"(?:(\d{4}) Results|Baseline (\d{4}))", c) for i, c in enumerate(r)}
            if any(hdr.values()) and any(c == "Unit" for c in r):
                years = {i: (m.group(1) or m.group(2)) for i, m in hdr.items() if m}
                tcol = next((i for i, c in enumerate(r) if c == "Future Target"), None)
                dcol = next((i for i, c in enumerate(r) if c.startswith("Target Completion")), None)
                continue
            if not years:
                continue
            label = re.sub(r"^[•�·\s]+", "", r[0])
            hit = next((m for m in SCORECARD_MAP if re.search(m[0], label, re.I)), None)
            vals = {y: parse_value(r[i]) for i, y in years.items() if i < len(r)}
            if not hit:
                if any(v for v, _, _ in vals.values()):
                    unmapped.add(label[:70])
                continue
            _, metric, mode, unit = hit
            for y, (v, note, prec) in vals.items():
                if v is not None:
                    out.append(rec(doc, page, "scorecard", metric, mode, "", y, "year", v, unit, note, prec))
            if tcol is not None and tcol < len(r):
                v, _, prec = parse_value(r[tcol])
                due = r[dcol] if dcol is not None and dcol < len(r) else ""
                if v is not None:
                    out.append(rec(doc, page, "scorecard", metric, mode, "target",
                                   f"edition {doc.meeting_date[:4]}", "target", v, unit,
                                   f"target completion: {due}" if due else "", prec))
    for u in sorted(unmapped):
        log.append(f"  unmapped scorecard row (numeric, skipped): {u}")
    return out


def table_segments(doc, page):
    """Split a page's blocks into tables: a new one starts at each 'Month' header."""
    segs = []
    for pg, rows in doc.blocks:
        if pg != page:
            continue
        for r in rows:
            first = next((c for c in r if c), "")
            if first.startswith("Month") or re.fullmatch(r"\d{4} Period", first):
                segs.append({"header": " | ".join(c for c in r if c), "rows": []})
            elif segs:
                if not month_period(first, 2000) and not segs[-1]["rows"] and not re.search(r"\d", first + "".join(r[1:2])):
                    segs[-1]["header"] += " | " + " | ".join(c for c in r if c)  # 2nd header row
                else:
                    segs[-1]["rows"].append([c for c in r if c])
    return segs


def monthly(doc, spec, log):
    """spec: pages, header (must contain), not_header, year (default), cols."""
    out = []
    for page in spec["pages"]:
        segs = [s for s in table_segments(doc, page)
                if spec["header"] in s["header"]
                and not (spec.get("not_header") and spec["not_header"] in s["header"])]
        if not segs:
            log.append(f"  NO MATCH p{page}: monthly table with header '{spec['header']}'")
            continue
        for row in segs[0]["rows"]:
            label, cells = row[0], row[1:]
            period, ptype = month_period(label, spec.get("year")), "month"
            if not period:
                period, ptype = total_period(label, spec.get("year"), log, f"p{page}")
            if not period:
                continue
            for col, cell in zip(spec["cols"], cells):
                if col is None:
                    continue
                v, note, prec = parse_value(cell)
                if v is not None:
                    metric, mode, breakdown, unit = col
                    out.append(rec(doc, page, "monthly", metric, mode, breakdown, period, ptype, v, unit,
                                   note, prec))
    return out


def chart_rows(doc, spec, log):
    """Lines like '2022 Actual 3.0 M 3.4 M ...' (Jan.. onward) on a given page."""
    out = []
    txt = doc.pages.get(spec["page"], "")
    tok = r"\$?\d+(?:\.\d+)? ?[Mk]"
    for m in re.finditer(rf"^(\d{{4}}) Actual ((?:{tok} ?)+)", txt, flags=re.M):
        year = int(m.group(1))
        for i, t in enumerate(re.findall(tok, m.group(2)), 1):
            v, _, prec = parse_value(t)
            out.append(rec(doc, spec["page"], "chart_rows", spec["metric"], spec["mode"], "",
                           f"{year}-{i:02d}", "month", v, spec["unit"],
                           "chart label, rounded to 0.1 " + ("M" if "M" in t else "k"), prec))
    if not out:
        log.append(f"  NO MATCH p{spec['page']}: chart rows for {spec['metric']}")
    return out


def sequence(doc, spec, log):
    """A line of N percentages after `anchor`, one per consecutive month from `start`."""
    txt = doc.pages.get(spec["page"], "")
    pos = txt.find(spec["anchor"])
    for line in txt[pos:].splitlines() if pos >= 0 else []:
        vals = re.findall(r"(\d+(?:\.\d+)?)%", line)
        if len(vals) >= spec["n"] - 1:
            break
    else:
        vals = []
    if len(vals) != spec["n"]:
        log.append(f"  NO MATCH p{spec['page']}: expected {spec['n']} labels for "
                   f"{spec['metric']}/{spec['mode']}, found {len(vals)}")
        return []
    y, m = map(int, spec["start"].split("-"))
    out = []
    for v in vals:
        out.append(rec(doc, spec["page"], "sequence", spec["metric"], spec["mode"], "",
                       f"{y}-{m:02d}", "month", float(v), "%", spec.get("note", ""),
                       parse_value(v + "%")[2]))
        y, m = (y + 1, 1) if m == 12 else (y, m + 1)
    return out


def narrative(doc, spec, log):
    """Regex over whitespace-normalised page text; each group -> one value."""
    for page in sorted(doc.pages):
        m = re.search(spec["re"], doc.flat(page), flags=re.I)
        if m:
            out = []
            for g, (metric, mode, breakdown, period, ptype, unit) in enumerate(spec["out"], 1):
                v, _, prec = parse_value(m.group(g))
                scale = spec.get("scale", 1)
                out.append(rec(doc, page, "narrative", metric, mode, breakdown, period, ptype,
                               round(v * scale, 4), unit, spec.get("note", ""), prec * abs(scale)))
            return out
    log.append(f"  NO MATCH: narrative /{spec['re'][:60]}.../")
    return []


# Open Ottawa "service & ridership KPIs" workbook: sheet -> [(column, metric, mode, unit, scale)].
# Ridership is published in millions and Para ridership in thousands.
SNAPSHOT_SHEETS = {
    "01_otp_conventional": [("regularity", "regularity_frequent_routes", "bus", "%", 1),
                            ("punctuality", "punctuality_less_frequent_routes", "bus", "%", 1)],
    "02_otp_para_transpo": [("overall", "para_on_time", "para", "%", 1)],
    "03_ridership": [("ridership", "ridership", "system", "linked trips", 1e6)],
    "04_ridership_paratranspo": [("paratranspo-ridership", "para_trips", "para", "trips", 1e3)],
    "05_otrain_service_delivery": [("o-train-line-1", "trips_delivered", "line1", "%", 1),
                                   ("o-train-line-2-4", "trips_delivered", "line2_4", "%", 1)],
    "06_bus_service_delivery": [("bus_service_delivery", "trips_delivered", "bus", "%", 1)],
}


def snapshots(log):
    """Monthly rows from every dated service_ridership_kpis_*.xlsx snapshot."""
    files = sorted(SNAPSHOTS.glob("service_ridership_kpis_*.xlsx"))
    if not files:
        log.append("[MISSING] no service_ridership_kpis snapshots in rt_data/kpi_snapshots/")
        return []
    try:
        import openpyxl
    except ImportError:
        log.append("[skip] KPI snapshots: openpyxl not installed (pip install openpyxl)")
        return []
    out = []
    for f in files:
        snap_date = re.search(r"(\d{4}-\d{2}-\d{2})", f.name).group(1)
        doc = type("Snap", (), {"meeting_date": snap_date, "key": f"rt_data/kpi_snapshots/{f.stem}"})
        wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
        found, sub = 0, []
        for sheet, cols in SNAPSHOT_SHEETS.items():
            if sheet not in wb.sheetnames:
                sub.append(f"  NO MATCH: sheet {sheet}")
                continue
            rows = list(wb[sheet].iter_rows(values_only=True))
            head = [str(h).strip() if h is not None else "" for h in rows[0]]
            for col, metric, mode, unit, scale in cols:
                if col not in head:
                    sub.append(f"  NO MATCH: {sheet} column '{col}'")
                    continue
                j = head.index(col)
                for r in rows[1:]:
                    m = re.fullmatch(r"([A-Za-z]{3}) (\d{2})", str(r[0] or "").strip())
                    if not m or m.group(1).lower() not in MONTHS or r[j] in (None, "", "-"):
                        continue
                    v, note, prec = parse_value(str(r[j]))
                    if v is None:
                        continue
                    period = f"20{m.group(2)}-{MONTHS[m.group(1).lower()]:02d}"
                    out.append(rec(doc, sheet, "snapshot", metric, mode, "", period, "month",
                                   int(round(v * scale)) if scale != 1 else v, unit, note, prec * scale))
                    found += 1
        wb.close()
        log.append(f"[{found:4d}] {doc.key}.xlsx (Open Ottawa snapshot)")
        log.extend(sub)
    return out


EXTRACTORS = {"scorecard": scorecard, "monthly": monthly, "chart_rows": chart_rows,
              "sequence": sequence, "narrative": narrative}

# ------------------------------------------------------------------------- specs

INFO_COLS = [("info_open_data_api_calls", "info", "", "count"),
             ("info_website_pageviews", "info", "", "count"),
             ("info_phone_schedule_calls", "info", "", "count"),
             ("info_sms_requests", "info", "", "count")]
EWT_COLS = [("excess_wait_time", "bus", "frequent", "min"),
            ("excess_wait_ratio", "bus", "frequent", "%"),
            ("excess_wait_time", "line1", "", "min"),
            ("excess_wait_ratio", "line1", "", "%")]
OTP_PERIOD_COLS = [("punctuality_less_frequent_routes", "bus", b, "%")
                   for b in ("weekday_am_peak", "weekday_pm_peak", "weekday_off_peak", "weekday_all_day")]
BUS_PROVISION_COLS = [("service_not_delivered", "bus", "", "%"),
                      None,  # frequent-service EWT, duplicate of the EWT table
                      ("punctuality_less_frequent_routes", "bus", "early", "%"),
                      ("punctuality_less_frequent_routes", "bus", "", "%"),
                      ("punctuality_less_frequent_routes", "bus", "late", "%")]
L1_PROVISION_COLS = [("service_not_delivered", "line1", "", "%"), None]

R2020 = "2021-04-21_Transit Commission/Report - OC Transpo Performance 2020 EN"
R2021H1 = "2021-10-20_Transit Commission/Document 1 - OC Transpo Performance Report for Period Ending June 2021 EN"
R2021 = ("2022-05-18_Transit Commission/04 - ACS2022-TSD-TS-0002 - OC Transpo Performance "
         "Report for the Period Ending December 2021 EN")
DECK2023 = "2023-09-14_Transit Commission/English Slides - OC Transpo update"
TTC0619 = "2020-02-19_Transit Commission/TTC 06-19 - LRT and Bus Ridership Data and Service Disruption Info"

SPECS = {
    # --- Twice-yearly performance reports (2021-2022): real monthly tables ----
    R2020: [
        {"kind": "monthly", "pages": [15, 16], "header": "API Calls", "year": 2020, "cols": INFO_COLS},
        {"kind": "monthly", "pages": [16, 17], "header": "O-Train Line 2", "year": 2020,
         "cols": [("service_availability", m, "", "%") for m in ("bus", "line1", "line2", "system")]},
        {"kind": "monthly", "pages": [19], "header": "Rapid and Frequent", "year": 2020, "cols": EWT_COLS},
        {"kind": "monthly", "pages": [22], "header": "Morning Peak", "year": 2020, "cols": OTP_PERIOD_COLS},
        {"kind": "monthly", "pages": [23], "header": "Service Not Delivered", "not_header": "Frequent",
         "year": 2020, "cols": L1_PROVISION_COLS},
        {"kind": "monthly", "pages": [23, 24], "header": "Frequent Service", "year": 2020,
         "cols": BUS_PROVISION_COLS},
        {"kind": "narrative", "re": r"100% " + r"([\d.]+)% " * 8 + r"80% 60% 40% 20% 0% Bus O-Train Line 1",
         "note": "Figure 12 chart labels (Bus, Line 1, Line 2, all; 2019 then 2020). Line 1 2019 = "
                 "Sep 14-Dec only; Line 2 2020 = Jan-May 2 only",
         "out": [("service_availability", "bus", "", "2019", "year", "%"),
                 ("service_availability", "bus", "", "2020", "year", "%"),
                 ("service_availability", "line1", "", "2019", "year", "%"),
                 ("service_availability", "line1", "", "2020", "year", "%"),
                 ("service_availability", "line2", "", "2019", "year", "%"),
                 ("service_availability", "line2", "", "2020", "year", "%"),
                 ("service_availability", "system", "", "2019", "year", "%"),
                 ("service_availability", "system", "", "2020", "year", "%")]},
        {"kind": "narrative", "re": r"injury rate in 2019 was ([\d.]+) injuries per million trips; this "
                                    r"increased 22 percent to ([\d.]+) in 2020",
         "out": [("customer_injury_rate", "system", "", "2019", "year", "injuries per million trips"),
                 ("customer_injury_rate", "system", "", "2020", "year", "injuries per million trips")]},
        {"kind": "narrative", "re": r"There were fewer injuries in 2020 than 2019: (\d+) in 2020, compared to (\d+) in 2019",
         "out": [("customer_injuries", "system", "", "2020", "year", "count"),
                 ("customer_injuries", "system", "", "2019", "year", "count")]},
        {"kind": "narrative", "re": r"ridership in the full calendar year 2020 was (\d+) percent lower than in 2019",
         "scale": -1, "out": [("ridership_change_vs_prior_year", "system", "", "2020", "year", "%")]},
        {"kind": "narrative", "re": r"Ottawa had (\d+) trips per capita in 2019, down slightly from (\d+) in 2018",
         "out": [("ridership_per_capita", "system", "", "2019", "year", "trips per resident"),
                 ("ridership_per_capita", "system", "", "2018", "year", "trips per resident")]},
        {"kind": "narrative", "re": r"increased from ([\d.]+) percent in 2019 to ([\d.]+) percent in 2020",
         "note": "within a 30-minute pick-up window",
         "out": [("para_on_time", "para", "", "2019", "year", "%"),
                 ("para_on_time", "para", "", "2020", "year", "%")]},
        {"kind": "narrative", "re": r"out of a total of ([\d,]+) requests\) that could not be accommodated in the calendar year 2020",
         "out": [("para_booking_requests", "para", "", "2020", "year", "count")]},
        {"kind": "narrative", "re": r"Overall, for 2020 elevator availability was ([\d.]+) percent",
         "out": [("elevator_availability", "system", "", "2020", "year", "%")]},
    ],
    R2021H1: [
        {"kind": "monthly", "pages": [7], "header": "API Calls", "cols": INFO_COLS},
        {"kind": "monthly", "pages": [8], "header": "Train and Bus",
         "cols": [("service_availability", m, "", "%") for m in ("bus", "line1", "system")]},
        {"kind": "monthly", "pages": [10], "header": "Rapid and Frequent", "cols": EWT_COLS},
        {"kind": "monthly", "pages": [13], "header": "Morning Peak", "cols": OTP_PERIOD_COLS},
        {"kind": "monthly", "pages": [13], "header": "Service Not Delivered", "not_header": "Frequent",
         "cols": L1_PROVISION_COLS},
        {"kind": "monthly", "pages": [14], "header": "Frequent Service", "cols": BUS_PROVISION_COLS},
    ],
    R2021: [
        {"kind": "monthly", "pages": [18], "header": "API Calls", "year": 2021, "cols": INFO_COLS},
        {"kind": "monthly", "pages": [23], "header": "Service Not Delivered", "year": 2021,
         "cols": L1_PROVISION_COLS},
        {"kind": "monthly", "pages": [24], "header": "Rapid and Frequent", "year": 2021, "cols": EWT_COLS},
        {"kind": "narrative", "re": r"2020 - Complaints per Trip " + r"([\d.]+) " * 11 + r"([\d.]+)",
         "note": "per 1,000 Para Transpo trips",
         "out": [("para_complaints_per_1000_trips", "para", "", f"2020-{i:02d}", "month", "per 1,000 trips")
                 for i in range(1, 13)]},
        {"kind": "narrative", "re": r"2021 - Complaints per Trip " + r"([\d.]+) " * 11 + r"([\d.]+)",
         "note": "per 1,000 Para Transpo trips",
         "out": [("para_complaints_per_1000_trips", "para", "", f"2021-{i:02d}", "month", "per 1,000 trips")
                 for i in range(1, 13)]},
        {"kind": "narrative", "re": r"2020 - Total Complaints " + r"(\d+) " * 11 + r"(\d+)",
         "out": [("para_complaints", "para", "", f"2020-{i:02d}", "month", "count") for i in range(1, 13)]},
        {"kind": "narrative", "re": r"2021 - Total Complaints " + r"(\d+) " * 11 + r"(\d+)",
         "out": [("para_complaints", "para", "", f"2021-{i:02d}", "month", "count") for i in range(1, 13)]},
        {"kind": "narrative", "re": r"bus service availability from (\d+) to (\d+) per cent",
         "out": [("service_availability", "bus", "", "2020", "year", "%"),
                 ("service_availability", "bus", "", "2021", "year", "%")]},
        {"kind": "narrative", "re": r"reduction from (\d+) per cent in 2020 to (\d+) per cent in 2021",
         "note": "Line 1 closed Aug 9-13 and Sep 19-Nov 12 2021 (derailments)",
         "out": [("service_availability", "line1", "", "2020", "year", "%"),
                 ("service_availability", "line1", "", "2021", "year", "%")]},
        {"kind": "narrative", "re": r"Excluding the derailment related closures, Line 1 service availability in 2021 was (\d+) per cent",
         "out": [("service_availability", "line1", "excl_derailment_closures", "2021", "year", "%")]},
    ],
    # --- 2023 update deck: chart labels survive in the text layer ------------
    DECK2023: [
        {"kind": "chart_rows", "page": 25, "metric": "ridership", "mode": "system", "unit": "linked trips"},
        {"kind": "chart_rows", "page": 20, "metric": "para_trips", "mode": "para", "unit": "trips"},
        {"kind": "chart_rows", "page": 26, "metric": "fare_revenue", "mode": "system", "unit": "CAD"},
        {"kind": "sequence", "page": 22, "anchor": "Rail Service Delivery", "n": 13, "start": "2022-08",
         "metric": "trips_delivered", "mode": "line1",
         "note": "chart label; Jul 2023 before the Jul 17 closure was 99.4%"},
        {"kind": "sequence", "page": 23, "anchor": "Bus Service Delivery", "n": 13, "start": "2022-08",
         "metric": "trips_delivered", "mode": "bus", "note": "chart label"},
    ],
    # --- 5-Year Roadmap scorecards: annual, one column per year --------------
    "2024-04-11_Transit Commission/Document 1 English - Transit Services 5-Year Roadmap Scorecard":
        [{"kind": "scorecard"}],
    "2025-04-10_Transit Committee/Document 1 - English - 5-Year Roadmap": [{"kind": "scorecard"}],
    "2026-05-14_Transit Committee/English Document 1 - Transit Services 5-Year Roadmap Scorecard 2025":
        [{"kind": "scorecard"}],
    # --- One-off memos / reports with a few quotable figures -----------------
    TTC0619: [
        {"kind": "narrative", "re": r"month of October 2018 was ([\d.]+) million",
         "scale": 1e6, "out": [("ridership", "system", "", "2018-10", "month", "linked trips")]},
        {"kind": "narrative", "re": r"for October 2019 was ([\d.]+) million customer trips",
         "scale": 1e6, "out": [("ridership", "system", "", "2019-10", "month", "linked trips")]},
        {"kind": "narrative", "re": r"In the first seven weeks of service on Line 1, ([\d.]+)% of scheduled service was delivered",
         "out": [("service_availability", "line1", "", "2019-09-14/2019-11-05", "range", "%")]},
    ],
    "2025-03-17_Transit Committee/English slides - OC Transpo Update Presentation": [
        {"kind": "narrative", "re": r"January 2025 ridership: ([\d,]+)",
         "note": "GM-update bullet under 'Full 7-day service on Lines 2 & 4'; Lines 2/4 only "
                 "(system ridership is ~5-6M/month)",
         "out": [("ridership", "line2_4", "", "2025-01", "month", "trips")]},
    ],
    "2025-09-11_Transit Committee/Improving Bus Service Reliability - English": [
        {"kind": "narrative", "re": r"Actual availability was (\d+) per cent in 2023 and improved to (\d+) per cent in 2024",
         "out": [("bus_operator_availability", "bus", "", "2023", "year", "%"),
                 ("bus_operator_availability", "bus", "", "2024", "year", "%")]},
        {"kind": "narrative", "re": r"reduced punctuality to (\d+) per cent for the month from (\d+) per cent in January and March",
         "note": "Feb 2025: two major snowfalls",
         "out": [("punctuality_less_frequent_routes", "bus", "", "2025-02", "month", "%"),
                 ("punctuality_less_frequent_routes", "bus", "", "2025-01", "month", "%")]},
        {"kind": "narrative", "re": r"Transit ridership has increased by (\d+) per cent between mid-2023",
         "out": [("ridership_change", "system", "", "2023-07/2025-06", "range", "%")]},
    ],
    "2026-06-11_Transit Committee/OCC-2026-08 - Inquiry - Plante - OC Transpo (BIL)": [
        {"kind": "narrative",
         "re": r"Early AM (\d+)% (\d+)% .*?AM Peak (\d+)% (\d+)% .*?Midday (\d+)% (\d+)% .*?"
               r"PM Peak (\d+)% (\d+)% .*?Evening (\d+)% (\d+)% .*?Night (\d+)% (\d+)%",
         "note": "share of weekday customer trips by time of day (pre-pandemic winter 2020 vs winter 2026)",
         "out": [("weekday_ridership_share", "system", b, f"{y}-winter", "season", "%")
                 for b in ("early_am", "am_peak", "midday", "pm_peak", "evening", "night")
                 for y in (2020, 2026)]},
    ],
}

# Copy the March value onto its own month (the Jan regex group says "January and March").
POST_COPIES = [(("punctuality_less_frequent_routes", "bus", "", "2025-01"), "2025-03")]

# English docs reviewed and deliberately not extracted (shows up in the log).
EXCLUDED = {
    "RMCO": "LRT regulatory compliance audit: findings counts, no service KPIs",
    "LRT Compliance Report": "LRT regulatory compliance audit: findings counts, no service KPIs",
    "RCMO Annual Report": "LRT regulatory compliance audit: findings counts, no service KPIs",
    "Regulatory Monitor and Compliance": "LRT regulatory compliance audit: findings counts, no service KPIs",
    "Five-Year English Roadmap": "2023 roadmap: targets only (results start in the 2024 scorecard)",
    "5-Year Roadmap.pdf": "cover report for the scorecard (Document 1 has the numbers)",
    "Annual Update  Transit Services 5-Year Roadmap": "cover report for the scorecard",
    "example semi-annual report": "template with illustrative data, not actuals",
    "Performance Measurement and Reporting": "defines the KPI set; no results",
    "Performance Report for Period Ending June 2021 EN.pdf": "cover report; Document 1 has the tables",
    "2022 Business Plan": "plan narrative, no KPI results",
    "OCC 21-19": "'approximately 98.5 per cent of trips are delivered' with no period; skipped",
    "Ridership Forecast": "qualitative forecast answer",
    "Update Presentation (EN)": "June 2025 deck: route 11 corridor review + fleet plan; no system KPIs",
}


# --------------------------------------------------------------------- pipeline

def is_french(name):
    return bool(re.search(r"\bFR\b|French|\(FR\)|_FR\b|FR\.docx", name))


def agree(a, b):
    """Same figure printed at different rounding (97.8 vs 98, 9.57M vs 9.6M)."""
    return abs(float(a["value"]) - float(b["value"])) <= 0.5 * max(a["_prec"], b["_prec"]) + 1e-9


def pick(obs):
    """Latest publication wins; an older value that agrees within rounding but is
    printed more precisely replaces it. Returns (best, values that disagree)."""
    obs = sorted(obs, key=lambda o: (o["meeting_date"], -METHOD_RANK[o["method"]]), reverse=True)
    best = obs[0]
    for o in obs[1:]:
        if agree(o, best) and o["_prec"] < best["_prec"]:
            best = o
    return best, sorted({o["value"] for o in obs if not agree(o, best)})


def main():
    if not DATA.exists():
        sys.exit(f"{DATA} not found -- run oc_transpo_kpi_scraper.py --extract first (local only)")
    log, obs = [], []
    with (DATA / "manifest.csv").open(encoding="utf-8") as f:
        manifest = [Path(r["local_path"].replace("\\", "/")) for r in csv.DictReader(f)]
    keys = [f"{p.parent.name}/{p.stem}" for p in manifest]

    log.append("== Per-document coverage ==")
    french = 0
    for key in sorted(set(keys) | set(SPECS)):
        name = key.split("/", 1)[1]
        if key not in SPECS:
            if is_french(name):
                french += 1
                continue
            reason = next((r for pat, r in EXCLUDED.items() if pat in name + ".pdf"), None)
            log.append(f"[skip] {key}: {reason or 'not KPI material (inquiry/memo/marketing/policy)'}")
            continue
        doc = Doc(key)
        if not doc.exists:
            log.append(f"[MISSING] {key}: no .txt/.tables.csv extract on disk")
            continue
        sub, found = [], []
        for spec in SPECS[key]:
            found += EXTRACTORS[spec["kind"]](doc, spec, sub)
        log.append(f"[{len(found):4d}] {key}")
        if not found:
            log.append("  WARNING: spec'd document yielded nothing")
        log.extend(sub)
        obs += found
    log.append(f"[skip] {french} French-language duplicates")
    obs += snapshots(log)

    for (metric, mode, bd, period), new_period in POST_COPIES:
        for o in [o for o in obs if (o["metric"], o["mode"], o["breakdown"], o["period"]) == (metric, mode, bd, period)
                  and o["method"] == "narrative"]:
            obs.append(dict(o, period=new_period))

    groups = defaultdict(list)
    for o in obs:
        groups[(o["metric"], o["mode"], o["breakdown"], o["period"], o["period_type"])].append(o)
    series, revisions = [], []
    for k in sorted(groups):
        best, alts = pick(groups[k])
        row = {f: best[f] for f in FIELDS}
        row["n_sources"] = len({o["source_file"] for o in groups[k]})
        row["other_values"] = ";".join(str(a) for a in alts)
        series.append(row)
        if alts:
            srcs = "; ".join(f"{o['value']} ({o['meeting_date']} p{o['page']} {o['method']})" for o in groups[k])
            revisions.append(f"  {'/'.join(x for x in k if x)}: kept {best['value']} -- {srcs}")

    log.append("\n== Revisions / conflicts between sources (kept value = kpi_series.csv) ==")
    log.extend(revisions or ["  none"])
    log.append("\n== Series coverage (kpi_series.csv) ==")
    cov = defaultdict(list)
    for r in series:
        cov[(r["metric"], r["mode"], r["breakdown"], r["period_type"])].append(r["period"])
    for (metric, mode, bd, pt), ps in sorted(cov.items()):
        log.append(f"  {metric:36s} {mode:8s} {bd or '-':26s} {pt:6s} {len(ps):3d}  {min(ps)} .. {max(ps)}")

    OUT.mkdir(exist_ok=True)
    with (OUT / "kpi_observations.csv").open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, FIELDS, extrasaction="ignore")
        w.writeheader()
        w.writerows(sorted(obs, key=lambda o: (o["metric"], o["mode"], o["breakdown"], o["period"], o["meeting_date"])))
    with (OUT / "kpi_series.csv").open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, FIELDS + ["n_sources", "other_values"])
        w.writeheader()
        w.writerows(series)
    (OUT / "extract_log.txt").write_text("\n".join(log) + "\n", encoding="utf-8")
    print(f"{len(obs)} observations -> {len(series)} series points; "
          f"{len(revisions)} revised/conflicting points; log: {OUT / 'extract_log.txt'}")


if __name__ == "__main__":
    main()
