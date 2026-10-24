"""Extract a reviewed housing snapshot. PDFs/workbook cached outside the repository.

Requires pymupdf and openpyxl. No collector inputs are modified. The 2025 memo
figures are manually reviewed against its public indexed text, explicitly noted.
"""
import hashlib
import json
import re
import tempfile
import urllib.request
from pathlib import Path

import openpyxl
import pymupdf

ROOT = Path(__file__).resolve().parents[1]
TEMP = Path(tempfile.gettempdir())
MEMO = 'https://ottawa.ca/en/city-hall/open-transparent-and-accountable-government/public-disclosure/memoranda-issued-members-council/memoranda-issued-community-and-social-services/memo-2025-housing-and-homelessness-plan-progress-update-june-26-2026'
SHEET = 'https://docs.google.com/spreadsheets/d/1VjBnjoMVAgX0e5R3HFgoDFWhNMo5RFX8bKt0PAddNhE/edit'
SOURCES = {
    'hna': {'title': 'City of Ottawa · 2024 Housing Needs Assessment', 'url': 'https://documents.ottawa.ca/sites/default/files/2024-HNA_EN.pdf', 'cache': 'ottawa-housing-hna.pdf', 'locator': 'Table 31, PDF p. 71; 2006–2023 completions', 'note': 'Use table values, not the contradictory prose beneath it: 2023 is 7,890 in the table; prose repeats 7,668, the 2022 figure. Components reconcile to totals.'},
    'ward': {'title': 'Councillor Jeff Leiper’s CMHC compilation', 'url': SHEET, 'download': SHEET.replace('/edit','/export?format=xlsx'), 'cache': 'ottawa-cmhc-ward.xlsx', 'locator': 'Completions by month; Starts roll-up; Under-construction roll-up; month-specific 2023–2026 tabs', 'note': 'Secondary manual compilation of CMHC census-subdivision releases, published by the councillor’s office. Notes worksheet identifies the CMHC CSD source. Not a direct CMHC download. Completion totals reconcile with the City HNA for 2020–2023; 2024/25 and 2026 YTD use the compilation.'},
    'report2023': {'title': 'City · 2023 Housing and Homelessness Progress Report', 'url': 'https://documents.ottawa.ca/sites/default/files/2023HHReport_EN.pdf', 'cache': 'housing2023.pdf', 'locator': 'PDF p. 7 (printed p. 4); completion projects on PDF p. 13', 'note': '49 affordable + 57 supportive = 106 completed units. PDF p. 13 separately mentions 40 units completed in late 2022 with move-ins in 2023; these are not added again.'},
    'report2024': {'title': 'City · 2024 Housing and Homelessness Progress Report', 'url': 'https://documents.ottawa.ca/sites/default/files/2024HHReport_EN.pdf', 'cache': 'housing2024.pdf', 'locator': 'PDF p. 10, 2024 by the numbers', 'note': '318 affordable and 75 supportive completed; 815 affordable and 112 supportive under construction at December 31. Year-end wait list and placements differ from HNA figures; retained separately.'},
    'memo2025': {'title': 'City · 2025 housing progress update, June 26, 2026', 'url': MEMO, 'locator': '2025 Housing and Homelessness Plan Progress Report, accomplishments list', 'note': 'Manually reviewed from indexed official memo text; direct automated opening returned 403. Reports 31 affordable + 48 supportive completed, 907 + 87 under construction at December 31, and 1,017 households housed from the wait list. No 2026 completion count in this source.'},
    'campaign_ad': {'title': 'Sutcliffe’s October 2022 campaign advertisement · Glebe Report', 'url': 'https://glebereport.ca/wp-content/uploads/2022/10/Glebe-Report-October-14-2022.pdf', 'cache': 'glebe2022.pdf', 'locator': 'PDF p. 34, campaign advertisement', 'note': 'Primary campaign evidence: 100,000 homes over the next ten years, including 10,000 affordable homes. No exact calendar baseline or interim milestones given.'},
    'campaign_news': {'title': 'September 16, 2022 housing platform · CityNews', 'url': 'https://ottawa.citynews.ca/2022/09/16/ottawa-mayoral-candidate-mark-sutcliffe-pledging-to-build-100k-new-homes-5829675/', 'locator': 'Housing pledge and community housing paragraphs', 'note': 'Contemporary report describes 100,000 new homes over ten years, including 10,000 community housing units and 1,000 community units per year. Terminology differs from the campaign advertisement’s affordable homes.'},
    'ward_page': {'title': 'Councillor’s explanation of the CMHC tracker', 'url': 'https://kitchissippiward.ca/2026/08/25/ottawa-housing-stats/', 'locator': 'CMHC Ottawa housing tracker; links to spreadsheet and CMHC CSD release', 'note': 'The downloadable workbook currently extends through August 2026. Its cutoff, rather than the article date, is used.'},
    'cmhc': {'title': 'CMHC · Starts, completions and under construction by census subdivisions', 'url': 'https://www.cmhc-schl.gc.ca/en/professionals/housing-markets-data-and-research/housing-data/data-tables/housing-market-data/starts-completions-under-construction-census-subdivisions', 'locator': 'Ottawa census subdivision, not Ottawa–Gatineau CMA', 'note': 'Upstream source identified by the ward compilation. Direct automated download was blocked; no claim of a fresh independent CMHC pull.'}
}


def fetch(source):
    p = TEMP / source['cache']
    if not p.exists():
        p.write_bytes(urllib.request.urlopen(source.get('download',source['url']), timeout=45).read())
    source['sha256'] = hashlib.sha256(p.read_bytes()).hexdigest()
    return p


if __name__ == '__main__':
    hna = pymupdf.open(fetch(SOURCES['hna']))
    text = hna[70].get_text()
    matches = re.findall(r'(?m)^(20\d{2})\n([\d,]+)\n([\d,]+)\n([\d,]+)\n([\d,]+)\n([\d,]+)\s*$', text)
    historic = []
    for year,*numbers in matches:
        values = [int(n.replace(',','')) for n in numbers]
        assert sum(values[:4]) == values[4]
        historic.append({'year': int(year), 'completions': values[4], 'source': 'hna', 'locator': 'Table 31, PDF p. 71', 'months':12})
    assert [r['year'] for r in historic] == list(range(2006,2024))
    workbook = openpyxl.load_workbook(fetch(SOURCES['ward']), data_only=True)
    completion_sheet = workbook['Completions by month']
    monthly = []
    for year in range(2020,2027):
        col = next(c.column for c in completion_sheet[1] if c.value == f'{year} monthly completions')
        ytd_col = col-1
        running = 0
        for month in range(1,13):
            cell = completion_sheet.cell(month+1,col)
            value = cell.value
            if value is None:
                continue
            assert int(value)==value and value>=0
            running += int(value)
            assert running == completion_sheet.cell(month+1,ytd_col).value, 'Monthly/YTD inconsistency'
            monthly.append({'period':f'{year}-{month:02}', 'completions':int(value), 'source':'ward', 'locator':f'Completions by month!{cell.coordinate}'})
    assert max(r['period'] for r in monthly)=='2026-08', 'Review cutoff and page labels before refreshing'
    for r in historic:
        if r['year']>=2020:
            assert r['completions']==sum(m['completions'] for m in monthly if m['period'].startswith(str(r['year'])))
    totals = historic[:]
    for year in [2024,2025,2026]:
        values = [r for r in monthly if r['period'].startswith(str(year))]
        totals.append({'year':year,'completions':sum(r['completions'] for r in values),'months':len(values),'source':'ward','locator':f'Completions by month, {year} monthly/YTD columns'})
    pipeline=[]
    for year in range(2023,2027):
        rows=[]
        s=workbook[f'By type rollup {year}']
        for month in range(1,13):
            col=2+(month-1)*3
            start,under,complete = (s.cell(7,c).value for c in [col,col+1,col+2])
            if complete is None:continue
            parts=[s.cell(r,col+2).value for r in range(3,7)]
            existing=next((r for r in monthly if r['period']==f'{year}-{month:02}'),None)
            if existing is None:
                assert complete==0 and all(v is None for v in parts), 'Unexpected unreported-month data'
                continue  # Formula zeros with blank components are placeholders, not observations.
            parts=[0 if isinstance(v,str) and v.strip() in ('-','–','—') else v for v in parts]
            assert all(isinstance(v,(int,float)) for v in parts), (year,month,parts,complete)
            assert sum(parts)==complete
            assert complete==existing['completions']
            rows.append((start,complete,under))
        pipeline.append({'year':year,'months':len(rows),'starts':int(sum(r[0] for r in rows)),'completions':int(sum(r[1] for r in rows)),'under_construction':int(rows[-1][2]),'stock_date':f'{year}-'+('12-31' if len(rows)==12 else '08-31'),'source':'ward','locator':f'By type rollup {year}, total row 7; 3-column blocks for starts/under construction/completions; components rows 3–6'})
    # Check primary report summary panels and retain their exact source pages.
    report2023=pymupdf.open(fetch(SOURCES['report2023']));report2024=pymupdf.open(fetch(SOURCES['report2024']))
    for number in ['49','57','12,447','1,186']:assert number in report2023[6].get_text()
    for number in ['318','75','14,721','1,077']:assert number in report2024[9].get_text()
    ad=pymupdf.open(fetch(SOURCES['campaign_ad']))
    assert '100,000' in ad[33].get_text() and '10,000' in ad[33].get_text()
    community=[{'year':2023,'affordable':49,'supportive':57,'completed':106,'source':'report2023','locator':'PDF p. 7','under_construction':831},
               {'year':2024,'affordable':318,'supportive':75,'completed':393,'source':'report2024','locator':'PDF p. 10','under_construction':927},
               {'year':2025,'affordable':31,'supportive':48,'completed':79,'source':'memo2025','locator':'Accomplishments list','under_construction':994}]
    data={'reviewed':'2026-10-24','geography':'City of Ottawa / Ottawa census subdivision; excludes neighbouring municipalities and Gatineau',
          'promise':{'homes':100000,'years':10,'community_homes':10000,'community_annual':1000,'analysis_start':'2023-01-01','analysis_start_note':'First full calendar year after taking office; analytical convention, not a verified campaign calendar baseline. Late 2022 excluded.'},
          'annual_completions':totals,'monthly_completions':monthly,'pipeline':pipeline,'affordable_supportive':community,
          'context':[{'year':2023,'year_end_waitlist':12447,'placements':1186,'source':'report2023'},{'year':2024,'year_end_waitlist':14721,'placements':1077,'source':'report2024'},{'year':2025,'year_end_waitlist':None,'placements':1017,'source':'memo2025'}],
          'reconciliation_notes':['HNA Table 31 has 7,890 completions in 2023; accompanying prose incorrectly repeats 7,668. Table chosen and checked against ward compilation.',
             'HNA Table 30 starts and its narrative disagree for several years. Starts here use only the ward CMHC CSD compilation; the HNA starts appendix is not spliced in.',
             '2024 HNA wait list: 15,140 on December 10; annual report: 14,721 on December 31. Different dates retained, not averaged or treated as explained revisions. HNA placements 1,089 versus annual report 1,077: reason not established. This section uses year-end annual report figures.',
             'Affordable/supportive report completions are not established as the exact scope of the campaign community/affordable pledge. No completion percentage for that pledge is assigned.',
             'Community figures are a subset/context of overall housing, not added to CMHC total completions; housing benefits and people housed are not new housing units.'],
          'sources':{k:{a:b for a,b in s.items() if a!='cache'} for k,s in SOURCES.items()}}
    (ROOT/'data/housing.json').write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(f'Extracted {len(totals)} annual/YTD completion observations, {len(monthly)} months and 3 affordable/supportive years.')
